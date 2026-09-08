# -*- coding: utf-8 -*-
"""
生成 3 个测试用的客户 API 文档样例（docs/测试案例/）
每个文档就是一份可直接上传到系统解析的 API 文档（.docx）：
  案例1 - 标准格式 API 文档（正常场景）
  案例2 - 格式不规范的 API 文档（异常场景，测容错）
  案例3 - 缺少关键字段的 API 文档（失败场景，测异常处理）
文档内容直接维护在本脚本中，修改后重新运行即可再生成
"""
import os
from docx import Document
from docx.shared import Pt, RGBColor

OUT_DIR = os.path.join(os.path.dirname(__file__), '..', 'docs', '测试案例')
os.makedirs(OUT_DIR, exist_ok=True)

def set_zh_font(run, size=10.5, bold=False, color=None):
    run.font.name = 'Calibri'
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = color
    r = run._element.rPr
    from docx.oxml.ns import qn
    rFonts = r.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = r.makeelement(qn('w:rFonts'), {})
        r.append(rFonts)
    rFonts.set(qn('w:eastAsia'), '微软雅黑')

def h1(doc, text):
    p = doc.add_paragraph()
    set_zh_font(p.add_run(text), 16, True, RGBColor(0x1F, 0x4E, 0x79))
    return p

def h2(doc, text):
    p = doc.add_paragraph()
    set_zh_font(p.add_run(text), 13, True, RGBColor(0x2E, 0x74, 0xB5))
    p.paragraph_format.space_before = Pt(10)
    return p

def h3(doc, text):
    p = doc.add_paragraph()
    set_zh_font(p.add_run(text), 11.5, True)
    p.paragraph_format.space_before = Pt(6)
    return p

def para(doc, text, bold=False):
    p = doc.add_paragraph()
    set_zh_font(p.add_run(text), 10.5, bold)
    return p

def bullets(doc, items):
    for it in items:
        p = doc.add_paragraph(style='List Bullet')
        set_zh_font(p.add_run(it))

def table(doc, headers, rows):
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = 'Table Grid'
    for j, htext in enumerate(headers):
        set_zh_font(t.rows[0].cells[j].paragraphs[0].add_run(htext), 10, True)
    for i, row in enumerate(rows):
        for j, cell in enumerate(row):
            set_zh_font(t.rows[i + 1].cells[j].paragraphs[0].add_run(cell), 10)
    return t

# ================= 案例 1：标准格式 API 文档（正常场景） =================
doc = Document()
h1(doc, '用户中心系统 API 接口文档')
para(doc, '版本：V1.2    发布日期：2026-08-15    维护部门：平台研发部')
para(doc, '本文档描述用户中心系统对外提供的接口服务，供合作方系统对接使用。所有接口均采用 HTTPS 协议，请求与响应报文格式为 JSON，字符编码 UTF-8。')

h2(doc, '5.1.1 用户登录接口')
h3(doc, '5.1.1.1 服务地址')
para(doc, 'POST /api/v1/auth/login')
h3(doc, '5.1.1.2 服务描述')
para(doc, '用户使用账号密码进行登录认证，认证通过后返回访问令牌（Token），后续接口调用需在请求头中携带该令牌。')
h3(doc, '5.1.1.3 输入参数')
table(doc, ['参数名', '类型', '必填', '说明'], [
    ['username', 'string', '是', '用户账号，长度 4-32 位'],
    ['password', 'string', '是', '用户密码，需使用 MD5 加密后传输'],
    ['captcha', 'string', '否', '图形验证码，连续登录失败 3 次后必填'],
])
h3(doc, '5.1.1.4 输出参数')
table(doc, ['参数名', '类型', '说明'], [
    ['code', 'int', '状态码，0 表示成功，非 0 表示失败'],
    ['message', 'string', '提示信息，失败时返回具体原因'],
    ['token', 'string', '访问令牌，有效期 2 小时'],
    ['expireAt', 'int', '令牌过期时间戳（秒）'],
])

h2(doc, '5.1.2 获取用户信息接口')
h3(doc, '5.1.2.1 服务地址')
para(doc, 'GET /api/v1/user/profile')
h3(doc, '5.1.2.2 服务描述')
para(doc, '获取当前登录用户的详细信息。请求头需携带有效的访问令牌（Authorization: Bearer {token}）。')
h3(doc, '5.1.2.3 输入参数')
table(doc, ['参数名', '类型', '必填', '说明'], [
    ['userId', 'int', '是', '用户 ID'],
])
h3(doc, '5.1.2.4 输出参数')
table(doc, ['参数名', '类型', '说明'], [
    ['code', 'int', '状态码，0 表示成功'],
    ['data', 'object', '用户信息对象'],
    ['data.id', 'int', '用户 ID'],
    ['data.username', 'string', '用户名'],
    ['data.email', 'string', '邮箱地址'],
    ['data.phone', 'string', '手机号码'],
])

h2(doc, '5.1.3 修改用户信息接口')
h3(doc, '5.1.3.1 服务地址')
para(doc, 'PUT /api/v1/user/profile')
h3(doc, '5.1.3.2 服务描述')
para(doc, '修改当前登录用户的基本信息。仅允许修改本人信息，邮箱与手机号修改后需重新验证。')
h3(doc, '5.1.3.3 输入参数')
table(doc, ['参数名', '类型', '必填', '说明'], [
    ['userId', 'int', '是', '用户 ID'],
    ['email', 'string', '否', '新邮箱地址'],
    ['phone', 'string', '否', '新手机号码'],
    ['nickname', 'string', '否', '昵称，长度 2-16 位'],
])
h3(doc, '5.1.3.4 输出参数')
table(doc, ['参数名', '类型', '说明'], [
    ['code', 'int', '状态码，0 表示成功'],
    ['message', 'string', '提示信息'],
])
doc.save(os.path.join(OUT_DIR, '测试案例1-标准API文档解析（正常场景）.docx'))

# ================= 案例 2：格式不规范的 API 文档（异常场景） =================
doc = Document()
h1(doc, '订单模块对接说明')
para(doc, '这是我们订单系统的接口说明，写得比较随意，对接的时候有问题随时联系张工。')

h2(doc, '订单提交')
para(doc, '这个接口用来提交订单，需要POST到 /order/submit')
para(doc, '参数包括：')
bullets(doc, [
    'orderId 订单编号（必填）',
    'items 商品列表（必填，数组类型）',
    'totalAmount 总金额（number类型）',
    'remark：备注（可选）',
])
para(doc, '返回的数据：')
para(doc, 'success - 是否成功（boolean）')
para(doc, 'orderId - 订单号')
para(doc, 'payUrl - 支付地址')

h2(doc, '查询订单状态')
para(doc, '请求方式: GET')
para(doc, '接口地址: /order/status')
para(doc, '入参:')
para(doc, 'orderId(string) - 订单ID')
para(doc, '出参:')
para(doc, '{')
para(doc, '  "status": "订单状态",')
para(doc, '  "createTime": "创建时间",')
para(doc, '  "updateTime": "更新时间"')
para(doc, '}')
doc.save(os.path.join(OUT_DIR, '测试案例2-格式不规范的文档（异常场景1）.docx'))

# ================= 案例 3：缺少关键字段的 API 文档（失败场景） =================
doc = Document()
h1(doc, '接口需求整理（草稿）')
para(doc, '以下是初步整理的接口需求，很多细节还没定，先发出来对齐方向。')

h2(doc, '数据同步接口')
para(doc, '这是一个用于同步数据的接口。')
para(doc, '参数：')
para(doc, 'dataId - 数据ID')
para(doc, 'timestamp - 时间戳')
para(doc, 'content - 内容')
para(doc, '返回：')
para(doc, 'result - 结果')

h2(doc, '发送通知')
para(doc, '输入：title, message, userId')
para(doc, '输出：success')

h2(doc, '第三个接口')
para(doc, '这个接口什么信息都不全。')
doc.save(os.path.join(OUT_DIR, '测试案例3-缺少关键字段的文档（异常场景2）.docx'))

print('已生成 3 个测试用 API 文档到:', os.path.abspath(OUT_DIR))
